"""Excel artifact writer.

Reproduces the layout of the legacy ``PandasWriterReader.writeExcel``
(``pandas.DataFrame.to_excel`` with a 1-based index): cell A1 empty, headers
in B1..P1, the 1-based row index in column A, and ragged columns padded with
blank cells. Written with openpyxl directly so pandas/numpy are not runtime
dependencies.
"""

from __future__ import annotations

import io
from typing import Any

from openpyxl import Workbook, load_workbook

from sheetcut.core.models import CutProgram


def write_xlsx(program: CutProgram) -> bytes:
    wb = Workbook()
    ws = wb.active
    assert ws is not None

    for col_idx, (header, _) in enumerate(program.columns(), start=2):
        ws.cell(row=1, column=col_idx, value=header)

    for row_idx, row in enumerate(program.rows(), start=2):
        ws.cell(row=row_idx, column=1, value=row_idx - 1)
        for col_idx, value in enumerate(row, start=2):
            if value is not None:
                ws.cell(row=row_idx, column=col_idx, value=value)

    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


def read_grid(source: bytes | str) -> list[list[Any]]:
    """Read an xlsx (ours or the legacy writer's) into a cell grid.

    Trailing all-``None`` rows/columns are trimmed so the two writers compare
    on content, not on how many empty cells they happened to touch.
    """
    if isinstance(source, bytes):
        wb = load_workbook(io.BytesIO(source), read_only=True, data_only=True)
    else:
        wb = load_workbook(source, read_only=True, data_only=True)
    ws = wb.active
    assert ws is not None
    grid = [list(row) for row in ws.iter_rows(values_only=True)]
    wb.close()

    while grid and all(v is None for v in grid[-1]):
        grid.pop()
    width = 0
    for row in grid:
        row_width = len(row)
        while row_width and row[row_width - 1] is None:
            row_width -= 1
        width = max(width, row_width)
    return [row[:width] + [None] * (width - len(row[:width])) for row in grid]
